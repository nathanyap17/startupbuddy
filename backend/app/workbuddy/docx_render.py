"""Render WorkBuddy's report_markdown into a styled DOCX (headings, paragraphs, bullets, tables, bold).
Server-side so the report is always downloadable from the web app, without WorkBuddy hunting for docx skills."""
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x23, 0x3E, 0x56)
SLATE = RGBColor(0x43, 0x47, 0x4D)
BOLD = re.compile(r"(\*\*[^*]+\*\*)")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
# Lone surrogates (broken emoji) and XML-illegal control characters break the DOCX XML.
INVALID = re.compile("[\ud800-\udfff\x00-\x08\x0b\x0c\x0e-\x1f]")


def _runs(par, text: str) -> None:
    text = LINK.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)
    for part in BOLD.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        else:
            par.add_run(part.replace("`", ""))


def _table(doc, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "") for c in r)]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Light Grid Accent 1"
    for i, r in enumerate(cells):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            _runs(cell.paragraphs[0], r[j] if j < len(r) else "")
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def render(markdown: str, title: str) -> bytes:
    markdown, title = INVALID.sub("", markdown), INVALID.sub("", title)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    for lvl, size in ((1, 20), (2, 15), (3, 12.5)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.color.rgb, h.font.size = NAVY, Pt(size)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = cover.add_run("StartupBuddy · From idea to first meeting.")
    r.font.color.rgb, r.font.size, r.bold = SLATE, Pt(9), True

    table_buf: list[str] = []
    first_h1 = True
    for line in markdown.replace("\r\n", "\n").split("\n") + [""]:
        if line.strip().startswith("|"):
            table_buf.append(line)
            continue
        if table_buf:
            _table(doc, table_buf)
            table_buf = []
        s = line.rstrip()
        if not s.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            level = min(len(m.group(1)), 3)
            text = m.group(2).strip()
            if level == 1 and first_h1:
                first_h1 = False
                doc.add_heading(text or title, level=0)
            else:
                doc.add_heading(text, level=level)
            continue
        m = re.match(r"^\s*[-*]\s+(.*)", s)
        if m:
            _runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", s)
        if m:
            _runs(doc.add_paragraph(style="List Number"), m.group(1))
            continue
        if re.fullmatch(r"\s*(-{3,}|\*{3,})\s*", s):
            continue
        _runs(doc.add_paragraph(), s.lstrip("> "))
    if first_h1:
        doc.paragraphs[0].insert_paragraph_before(title, style="Title")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
